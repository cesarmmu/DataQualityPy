"""Exportación del dataset limpio de DataQualityPy V6."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from shutil import copy2
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from src.loader import ErrorCargaArchivo, obtener_metadatos_csv


class ErrorExportacionLimpieza(Exception):
    """Error producido al exportar un dataset limpio."""


@dataclass(frozen=True)
class ResultadoExportacion:
    """Describe el archivo limpio generado."""

    formato: str
    ruta: Path
    nombre_archivo: str
    filas: int
    columnas: int
    hoja: str | None = None
    delimitador: str | None = None
    codificacion: str | None = None

    def a_dict(self) -> dict[str, Any]:
        """Devuelve una representación serializable del resultado."""
        return {
            "formato": self.formato,
            "ruta": str(self.ruta),
            "nombre_archivo": self.nombre_archivo,
            "filas": self.filas,
            "columnas": self.columnas,
            "hoja": self.hoja,
            "delimitador": self.delimitador,
            "codificacion": self.codificacion,
        }


def generar_nombre_limpio(nombre_original: str | Path) -> str:
    """Genera el nombre del archivo limpio."""
    ruta = Path(nombre_original)
    extension = ruta.suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise ErrorExportacionLimpieza(
            f"Formato no compatible para exportación: '{ruta.suffix}'."
        )
    return f"{ruta.stem}_clean{extension}"


def _preparar_ruta_salida(
    ruta_original: Path,
    ruta_salida: str | Path | None,
) -> Path:
    """Resuelve la salida e impide sobrescribir el original."""
    salida = (
        ruta_original.with_name(generar_nombre_limpio(ruta_original.name))
        if ruta_salida is None
        else Path(ruta_salida)
    )
    if salida.resolve(strict=False) == ruta_original.resolve(strict=False):
        raise ErrorExportacionLimpieza(
            "El archivo limpio no puede sobrescribir el archivo original."
        )
    salida.parent.mkdir(parents=True, exist_ok=True)
    return salida


def exportar_csv(
    datos_limpios: pd.DataFrame,
    ruta_original: str | Path,
    ruta_salida: str | Path | None = None,
) -> ResultadoExportacion:
    """Exporta CSV limpio conservando el delimitador detectado."""
    original = Path(ruta_original)
    if original.suffix.lower() != ".csv":
        raise ErrorExportacionLimpieza(
            "La exportación CSV requiere un archivo original .csv."
        )

    salida = _preparar_ruta_salida(original, ruta_salida)
    try:
        metadatos = obtener_metadatos_csv(original)
    except ErrorCargaArchivo as error:
        raise ErrorExportacionLimpieza(str(error)) from error

    delimitador = metadatos["delimitador"] or ","
    codificacion = "utf-8-sig"

    try:
        datos_limpios.to_csv(
            salida,
            index=False,
            sep=delimitador,
            encoding=codificacion,
        )
    except (OSError, UnicodeError, ValueError) as error:
        raise ErrorExportacionLimpieza(
            f"No se ha podido generar el CSV limpio '{salida}'."
        ) from error

    return ResultadoExportacion(
        "csv", salida, salida.name, len(datos_limpios),
        len(datos_limpios.columns),
        delimitador=delimitador,
        codificacion=codificacion,
    )


def _valor_excel(valor: Any) -> Any:
    """Convierte valores pandas/numpy a valores compatibles con Excel."""
    if pd.isna(valor):
        return None
    if hasattr(valor, "item"):
        try:
            return valor.item()
        except (ValueError, AttributeError):
            pass
    return valor


def _sustituir_datos_hoja(
    hoja: Any,
    datos_limpios: pd.DataFrame,
) -> None:
    """Vacía valores previos y escribe el DataFrame desde A1."""
    for fila in hoja.iter_rows():
        for celda in fila:
            celda.value = None

    for columna, nombre in enumerate(datos_limpios.columns, start=1):
        hoja.cell(1, columna, str(nombre))

    for fila, valores in enumerate(
        datos_limpios.itertuples(index=False, name=None),
        start=2,
    ):
        for columna, valor in enumerate(valores, start=1):
            hoja.cell(fila, columna, _valor_excel(valor))


def exportar_excel(
    datos_limpios: pd.DataFrame,
    ruta_original: str | Path,
    hoja: str,
    ruta_salida: str | Path | None = None,
) -> ResultadoExportacion:
    """Copia el XLSX y sustituye solo los datos de la hoja auditada."""
    original = Path(ruta_original)
    if original.suffix.lower() != ".xlsx":
        raise ErrorExportacionLimpieza(
            "La exportación Excel requiere un archivo original .xlsx."
        )
    if not hoja or not hoja.strip():
        raise ErrorExportacionLimpieza(
            "Debe indicarse la hoja auditada."
        )

    salida = _preparar_ruta_salida(original, ruta_salida)
    try:
        copy2(original, salida)
        libro = load_workbook(salida)
        if hoja not in libro.sheetnames:
            libro.close()
            raise ErrorExportacionLimpieza(
                f"La hoja auditada '{hoja}' no existe en el libro."
            )
        _sustituir_datos_hoja(libro[hoja], datos_limpios)
        libro.save(salida)
        libro.close()
    except ErrorExportacionLimpieza:
        salida.unlink(missing_ok=True)
        raise
    except Exception as error:
        salida.unlink(missing_ok=True)
        raise ErrorExportacionLimpieza(
            f"No se ha podido generar el Excel limpio '{salida}'."
        ) from error

    return ResultadoExportacion(
        "xlsx", salida, salida.name, len(datos_limpios),
        len(datos_limpios.columns), hoja=hoja,
    )


def exportar_dataset_limpio(
    datos_limpios: pd.DataFrame,
    ruta_original: str | Path,
    *,
    hoja: str | None = None,
    ruta_salida: str | Path | None = None,
) -> ResultadoExportacion:
    """Exporta el dataset limpio según el formato original."""
    original = Path(ruta_original)
    if not original.is_file():
        raise ErrorExportacionLimpieza(
            f"No se ha encontrado el archivo original: '{original}'."
        )

    if original.suffix.lower() == ".csv":
        return exportar_csv(datos_limpios, original, ruta_salida)

    if original.suffix.lower() == ".xlsx":
        if hoja is None:
            raise ErrorExportacionLimpieza(
                "Debe indicarse la hoja auditada para exportar un XLSX."
            )
        return exportar_excel(
            datos_limpios, original, hoja, ruta_salida
        )

    raise ErrorExportacionLimpieza(
        f"Formato no compatible: '{original.suffix}'."
    )

